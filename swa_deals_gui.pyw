"""
GUI for fetching live Southwest Airlines one-way flight deals from a chosen
home airport (BWI by default) and exporting them as a PDF table or an Excel
workbook.

Data is fetched live from southwest.com (see fetch_station_directory and
fetch_fares) the same way build_bwi_deals.py does it for BWI specifically.

See CHANGELOG.md for the version history of this file.
"""
import json
import math
import os
import queue
import re
import sqlite3
import sys
import threading
import time
import tkinter as tk
import urllib.error
import urllib.parse
import urllib.request
from contextlib import closing
from datetime import datetime, timedelta
from tkinter import messagebox, ttk

# Shown in the Help > About box — bump APP_VERSION alongside CHANGELOG.md entries.
APP_NAME = "SWA Flight Checker"
APP_VERSION = "1.13.0"
APP_DESCRIPTION = (
    "Fetches live one-way Southwest Airlines flight deals departing your "
    "chosen home airport directly from southwest.com, with distance and "
    "miles-per-dollar stats, sortable results, PDF/XLSX export, and "
    "optional auto-refresh."
)
AUTHOR_NAME = "Mike Jacobson"
AUTHOR_EMAIL = "coachmikej@gmail.com"

# Treeview column -> index into a display row tuple (rank, city, state, code, dist, date, price, mpd).
# Keep this in sync with the tuple shape returned by fetch_deals().
COLUMN_INDEX = {"rank": 0, "city": 1, "state": 2, "airport": 3, "distance": 4, "departing": 5, "price": 6, "mpd": 7}
# Columns sorted as case-insensitive text; everything else sorts numerically/chronologically.
TEXT_COLUMNS = {"city", "state", "airport"}

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

DEFAULT_ORIGIN = "BWI"
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"}
STATIONS_DATA_URL = "https://www.southwest.com/swa-ui/bootstrap/dotcom-landing-page/1/data.js"
LFMA_DATA_URL = "https://www.southwest.com/swa-resources/scripts/landing-pages/lfma.js"

# When packaged with PyInstaller --onefile, __file__ points into a temp folder
# that's deleted on exit, so save next to the .exe itself instead.
if getattr(sys, "frozen", False):
    OUTPUT_DIR = os.path.dirname(sys.executable)
else:
    OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))


def output_path(origin_code, extension):
    """Per-origin output filename, e.g. BWI_Flight_Deals.pdf / ORD_Flight_Deals.xlsx
    — keeps exports for different home airports from overwriting each other."""
    return os.path.join(OUTPUT_DIR, f"{origin_code}_Flight_Deals.{extension}")


# Local price history (SQLite). Lives next to the script/exe like the exports.
HISTORY_DB_PATH = os.path.join(OUTPUT_DIR, "price_history.db")
HISTORY_TIME_FORMAT = "%Y-%m-%d %H:%M:%S"


def _open_history_db():
    conn = sqlite3.connect(HISTORY_DB_PATH)
    conn.execute(
        """CREATE TABLE IF NOT EXISTS price_history (
               id INTEGER PRIMARY KEY AUTOINCREMENT,
               origin TEXT NOT NULL,
               dest TEXT NOT NULL,
               recorded_at TEXT NOT NULL,
               price INTEGER NOT NULL,
               depart_date TEXT
           )"""
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_route ON price_history (origin, dest, id)")
    return conn


def record_prices(origin_code, rows, when=None):
    """Append each route's price to the history, but only if it differs from the
    last price stored for that route (or the route has no history yet), so the
    table grows with price *changes*, not with every fetch. Returns how many
    rows were written. `rows` are fetch_deals() rows."""
    stamp = (when or datetime.now()).strftime(HISTORY_TIME_FORMAT)
    with closing(_open_history_db()) as conn:
        last = dict(conn.execute(
            """SELECT p.dest, p.price FROM price_history p
               JOIN (SELECT dest, MAX(id) AS last_id FROM price_history
                     WHERE origin = ? GROUP BY dest) m ON p.id = m.last_id""",
            (origin_code,),
        ))
        changed = [
            (origin_code, code, stamp, price, date_obj.isoformat())
            for _rank, _city, _state, code, _dist, date_obj, price, _mpd in rows
            if last.get(code) != price
        ]
        with conn:  # one transaction
            conn.executemany(
                "INSERT INTO price_history (origin, dest, recorded_at, price, depart_date) VALUES (?, ?, ?, ?, ?)",
                changed,
            )
    return len(changed)


def load_price_history(origin_code, dest_code):
    """Oldest-first list of (recorded_at: datetime, price, depart_date: str|None)."""
    with closing(_open_history_db()) as conn:
        found = conn.execute(
            "SELECT recorded_at, price, depart_date FROM price_history "
            "WHERE origin = ? AND dest = ? ORDER BY id",
            (origin_code, dest_code),
        ).fetchall()
    return [(datetime.strptime(at, HISTORY_TIME_FORMAT), price, dep) for at, price, dep in found]


# --- Settings (settings.json): Telegram bot config + price watches -----------
# Holds the bot token, which is a secret, so it's git-ignored. Kept next to the
# script/exe like the exports and history database.
SETTINGS_PATH = os.path.join(OUTPUT_DIR, "settings.json")


def default_settings():
    return {"telegram": {"bot_token": "", "chat_id": ""}, "watches": []}


def load_settings():
    """Returns (settings, warning). A missing file is normal (fresh install). A
    file that isn't valid JSON is moved aside to settings.json.bad rather than
    silently overwritten, and `warning` says so."""
    settings = default_settings()
    try:
        with open(SETTINGS_PATH, encoding="utf-8") as f:
            data = json.load(f)
    except FileNotFoundError:
        return settings, None
    except (json.JSONDecodeError, UnicodeDecodeError, OSError) as exc:
        try:
            os.replace(SETTINGS_PATH, SETTINGS_PATH + ".bad")
            moved = " (the old file was kept as settings.json.bad)"
        except OSError:
            moved = ""
        return settings, f"Couldn't read settings.json ({exc}); starting with empty settings{moved}."

    if not isinstance(data, dict):
        return settings, "settings.json isn't in the expected format; starting with empty settings."
    telegram = data.get("telegram")
    if isinstance(telegram, dict):
        settings["telegram"]["bot_token"] = str(telegram.get("bot_token") or "")
        settings["telegram"]["chat_id"] = str(telegram.get("chat_id") or "")
    for watch in data.get("watches") or []:
        # Skip malformed entries instead of failing the whole file.
        if (isinstance(watch, dict) and watch.get("origin") and watch.get("dest")
                and isinstance(watch.get("target"), (int, float)) and watch["target"] > 0):
            settings["watches"].append({
                "origin": str(watch["origin"]),
                "dest": str(watch["dest"]),
                "city": str(watch.get("city") or ""),
                "state": str(watch.get("state") or ""),
                "target": watch["target"],
                "created": watch.get("created"),
                "last_alert_price": watch.get("last_alert_price"),
                "last_alert_at": watch.get("last_alert_at"),
            })
    return settings, None


def save_settings(settings):
    """Write via a temp file + rename so a crash mid-write can't leave a
    half-written (unparseable) settings.json."""
    tmp_path = SETTINGS_PATH + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(settings, f, indent=2)
    os.replace(tmp_path, SETTINGS_PATH)


def fmt_money(amount):
    return f"${amount:,.0f}" if float(amount).is_integer() else f"${amount:,.2f}"


# --- Telegram -------------------------------------------------------------------
TELEGRAM_API_URL = "https://api.telegram.org/bot{token}/{method}"


class TelegramError(Exception):
    """A Telegram call failed; the message is safe to show (never contains the token)."""


def telegram_call(token, method, params=None, timeout=10):
    """POST to the Telegram Bot API and return the `result` field. Telegram
    explains failures in the JSON body ({"ok": false, "description": ...}), e.g.
    "Unauthorized" for a bad token or "chat not found" for a wrong chat ID."""
    url = TELEGRAM_API_URL.format(token=token, method=method)
    request = urllib.request.Request(url, data=urllib.parse.urlencode(params or {}).encode("utf-8"))
    try:
        with urllib.request.urlopen(request, timeout=timeout) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        try:
            description = json.loads(exc.read().decode("utf-8")).get("description")
        except (ValueError, OSError):
            description = None
        raise TelegramError(description or f"Telegram returned HTTP {exc.code}") from None
    except NETWORK_ERRORS as exc:
        # `from None` + our own message: keeps the request URL (which contains the token) out of any traceback.
        reason = getattr(exc, "reason", exc)
        raise TelegramError(f"Couldn't reach Telegram ({reason})") from None
    except ValueError:
        raise TelegramError("Telegram sent a reply that couldn't be read") from None
    if not body.get("ok"):
        raise TelegramError(body.get("description") or "Telegram rejected the request")
    return body.get("result")


def send_telegram(token, chat_id, text):
    telegram_call(token, "sendMessage", {"chat_id": chat_id, "text": text})


TELEGRAM_CHAT_KINDS = {"private": "private chat", "group": "group", "supergroup": "group", "channel": "channel"}


def detect_telegram_chats(token):
    """Every chat the bot has seen recently, newest first, as (chat_id, name, kind) —
    private chats, groups, and channels. That includes the "bot was added" event
    Telegram sends when you add the bot to a channel/group, so a freshly added
    channel shows up even before anyone posts in it."""
    updates = telegram_call(token, "getUpdates", {"limit": 100}) or []
    found = {}
    for update in reversed(updates):
        source = (update.get("message") or update.get("edited_message") or update.get("channel_post")
                  or update.get("edited_channel_post") or update.get("my_chat_member"))
        chat = (source or {}).get("chat") or {}
        if chat.get("id") is None or str(chat["id"]) in found:
            continue
        name = chat.get("title") or " ".join(
            part for part in (chat.get("first_name"), chat.get("last_name")) if part
        ) or chat.get("username") or "chat"
        found[str(chat["id"])] = (str(chat["id"]), name, TELEGRAM_CHAT_KINDS.get(chat.get("type"), "chat"))
    if not found:
        raise TelegramError(
            "No chats found. Send your bot a message (e.g. \"hi\"). For a channel or group, add the "
            "bot to it (channels: as an administrator who can post messages) and post a message there. "
            "Then click Detect again."
        )
    return list(found.values())


def booking_url(origin, dest, depart_date):
    """Deep link to Southwest's booking page for this route/date — the same URL
    format the flight-deals page's own \"Book now\" links use."""
    query = urllib.parse.urlencode({
        "promoCode": "", "adultPassengersCount": 1, "lapInfantPassengersCount": 0,
        "departureDate": depart_date.isoformat(), "destinationAirportCode": dest,
        "fareType": "USD", "originationAirportCode": origin, "passengerType": "ADULT",
        "returnDate": "", "tripType": "oneway", "departureTimeOfDay": "ALL_DAY",
        "reset": "true", "returnTimeOfDay": "ALL_DAY",
    })
    return f"https://www.southwest.com/air/booking/select-depart.html?{query}"


def format_alert(watch, row):
    """Plain-text Telegram message (no Markdown, so city names with ()_ can't break formatting)."""
    _rank, city, state, code, _dist, depart_date, price, _mpd = row
    return (
        f"✈️ SWA price alert\n"
        f"{watch['origin']} → {code}  ({city}, {state})\n"
        f"Now {fmt_money(price)} one-way — at or below your {fmt_money(watch['target'])} watch\n"
        f"Departing {fmt_date(depart_date)}\n"
        f"{booking_url(watch['origin'], code, depart_date)}"
    )


# Floor on the auto-refresh interval so a typo (or "0") can't hammer southwest.com.
MIN_AUTO_REFRESH_MINUTES = 1


FETCH_ATTEMPTS = 4
FETCH_RETRY_DELAY_SECONDS = 2

NETWORK_ERRORS = (urllib.error.URLError, TimeoutError, ConnectionError)
# How long the GUI waits before quietly retrying after a connection failure.
# Right after standby or login, Wi-Fi/VPN/DNS can take 20-60s to come back, which
# is longer than fetch()'s own short retries cover.
NETWORK_RETRY_SECONDS = 30


def is_network_error(exc):
    """True for connection-level failures (DNS, no route, timeout) as opposed to
    the server answering with an HTTP error or us failing to parse its data."""
    return isinstance(exc, NETWORK_ERRORS) and not isinstance(exc, urllib.error.HTTPError)


def fetch(url):
    """GET a URL, retrying connection-level failures (DNS lookup errors like
    getaddrinfo 11001, dropped Wi-Fi/VPN, timeouts, resets). Those are
    usually momentary, so a short pause and retry clears them. HTTP errors
    (404, 403, ...) mean the server answered, so they're raised right away."""
    req = urllib.request.Request(url, headers=HEADERS)
    for attempt in range(1, FETCH_ATTEMPTS + 1):
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                return resp.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError:
            raise
        except NETWORK_ERRORS:
            if attempt == FETCH_ATTEMPTS:
                raise
            time.sleep(FETCH_RETRY_DELAY_SECONDS * attempt)  # 2s, 4s, 6s


def fetch_station_directory():
    """Southwest ships airport reference data (lat/long, display name) as
    one module inside a large browserify bundle (many unrelated modules in
    one file). We don't need a JS engine to get at it — the module body is
    `module.exports = {...}` with plain JSON, so we locate it by module
    name and slice out the JSON between the `module.exports = ` marker and
    the bundle's closing `},{}]` for that module, then json.loads it
    directly. Returns code -> {"lat", "lon", "name"}."""
    js = fetch(STATIONS_DATA_URL)
    marker = '"swa-bootstrap-dotcom-landing-page/air/stations-data"'
    idx = js.index(marker)
    start = js.index("module.exports = ", idx) + len("module.exports = ")
    end = js.index("\n},{}]", start)
    stations = json.loads(js[start:end].rstrip().rstrip(";"))
    return {
        s["id"]: {
            "lat": float(s["latitude"]),
            "lon": float(s["longitude"]),
            "name": s.get("displayName") or s["id"],
        }
        for s in stations
        # A handful of entries (e.g. code-share-only stations) have no
        # coordinates published; skip them rather than crash on float(None).
        if s.get("latitude") is not None and s.get("longitude") is not None
    }


def fetch_fares(origin_code):
    """lfma.js ("Low Fare Market Ad") is the data file behind Southwest's
    own flight-deals page — it's regenerated periodically with the current
    live fares baked in as plain JS assignments like
    priceHash['BWI-MCO']='183'. We regex those straight out instead of
    executing the file, which also means we get every route Southwest has
    fare data for, not just the subset its own widget chooses to render.

    Returns (fares, origin_label): fares is a list of
    (dest_name, dest_code, departing_date, price) for routes departing
    origin_code, and origin_label is origin_code's own "City, ST" name
    (pulled from the same stationNameHash lfma.js uses for destinations)."""
    js = fetch(LFMA_DATA_URL)
    prices = dict(re.findall(r"priceHash\['([A-Z]+-[A-Z]+)'\]='(\d+)'", js))
    dates = dict(re.findall(r"departDateHash\['([A-Z]+-[A-Z]+)'\]='([\d/]+)'", js))
    names = dict(re.findall(r"stationNameHash\['([A-Z]+)'\]=\"(.*?)\";", js))

    fares = []
    for route, price in prices.items():
        origin, dest = route.split("-")
        if origin != origin_code or route not in dates or dest not in names:
            continue
        # Keep the year (not just "M/D") — needed so dates that roll over
        # into the next calendar year still sort chronologically later.
        date_obj = datetime.strptime(dates[route], "%m/%d/%Y").date()
        fares.append((names[dest], dest, date_obj, int(price)))
    return fares, names.get(origin_code, origin_code)


def fetch_origin_airports():
    """Airport codes Southwest treats as an origin market (i.e. you can pick
    them as a departure city), paired with a human-readable "City, ST"
    name, sorted alphabetically by that name. The code list comes from
    lfma.js's originArray; the names come from fetch_station_directory()."""
    js = fetch(LFMA_DATA_URL)
    match = re.search(r"originArray=\[(.*?)\];", js)
    if not match:
        raise RuntimeError("Could not find originArray in Southwest's data feed — its page structure may have changed.")
    codes = re.findall(r"'([A-Z]+)'", match.group(1))

    directory = fetch_station_directory()
    airports = [(code, directory[code]["name"]) for code in codes if code in directory]
    airports.sort(key=lambda pair: pair[1])
    return airports


def haversine_miles(lat1, lon1, lat2, lon2):
    """Great-circle (as-the-crow-flies) distance in miles, not flown route
    miles — actual flights zigzag around weather/traffic and are longer."""
    R = 3958.8  # mean Earth radius, miles
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def fmt_date(d):
    return f"{d.month}/{d.day:02d}"


def split_destination(name):
    """Southwest's station names are "City, ST" (domestic) or "City, Country"
    (international) — split on the last comma into (city, state_or_country)."""
    city, _sep, region = name.rpartition(",")
    return (city.strip(), region.strip()) if _sep else (name.strip(), "")


def fetch_deals(origin_code):
    """Returns (origin_label, rows) for flights departing origin_code.
    rows is a list of (rank, city, state, airport, distance_mi,
    departing_date, price, miles_per_dollar), sorted by price ascending.
    rank is each row's fixed price rank (1 = cheapest) and stays attached
    to the row even if it's later displayed in a different sort order."""
    directory = fetch_station_directory()
    if origin_code not in directory:
        raise RuntimeError(f"{origin_code!r} isn't a station Southwest publishes coordinates for.")

    fares, origin_label = fetch_fares(origin_code)
    if not fares:
        raise RuntimeError(
            f"No fares found for {origin_code} — it may not be a Southwest "
            f"origin market, or southwest.com's page structure may have changed."
        )

    origin = directory[origin_code]
    rows = []
    for dest_name, code, date_obj, price in fares:
        if code not in directory:
            continue
        dest = directory[code]
        dist = round(haversine_miles(origin["lat"], origin["lon"], dest["lat"], dest["lon"]))
        city, state = split_destination(dest_name)
        rows.append((city, state, code, dist, date_obj, price, dist / price))

    rows.sort(key=lambda r: r[5])
    ranked = [
        (rank, city, state, code, dist, date_obj, price, mpd)
        for rank, (city, state, code, dist, date_obj, price, mpd) in enumerate(rows, start=1)
    ]
    return origin_label, ranked


def export_pdf(rows, path, origin_code, origin_label):
    doc = SimpleDocTemplate(
        path,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch, rightMargin=0.5 * inch,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(
        f"Southwest Airlines &mdash; One-Way Flight Deals from {origin_label} ({origin_code})",
        styles["Title"],
    ))
    captured = datetime.now().strftime("%m/%d/%Y %H:%M")
    story.append(Paragraph(
        f"Live data fetched from southwest.com &mdash; captured {captured}. "
        f"Distances are great-circle (straight-line) miles between airports, not flown route miles.",
        styles["Normal"],
    ))
    story.append(Spacer(1, 0.2 * inch))

    table_data = [["Rank", "City", "State", "Airport", f"Distance from {origin_code} (mi)", "Departing", "One-Way Price ($)", "Miles per $"]]
    for rank, city, state, code, dist, date_obj, price, mpd in rows:
        table_data.append([str(rank), city, state, code, f"{dist:,}", fmt_date(date_obj), f"${price:,}", f"{mpd:.2f}"])

    table = Table(
        table_data,
        colWidths=[0.5 * inch, 1.9 * inch, 0.6 * inch, 0.7 * inch, 1.3 * inch, 0.9 * inch, 1.0 * inch, 0.8 * inch],
        repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004C97")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (3, -1), "CENTER"),
        ("ALIGN", (4, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)

    story.append(Spacer(1, 0.25 * inch))
    prices = [r[6] for r in rows]
    dists = [r[4] for r in rows]
    mpds = [r[7] for r in rows]
    summary = (
        f"Destinations: {len(rows)} &nbsp;|&nbsp; "
        f"Cheapest: ${min(prices):,} &nbsp;|&nbsp; Most expensive: ${max(prices):,} &nbsp;|&nbsp; "
        f"Average fare: ${sum(prices) / len(prices):,.2f} &nbsp;|&nbsp; "
        f"Average distance: {sum(dists) / len(dists):,.0f} mi &nbsp;|&nbsp; "
        f"Average miles per $: {sum(mpds) / len(mpds):,.2f}"
    )
    story.append(Paragraph(summary, styles["Normal"]))

    doc.build(story)


def export_xlsx(rows, path, origin_code, origin_label):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{origin_code} Flight Deals"

    font_name = "Arial"
    header_font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="004C97", end_color="004C97", fill_type="solid")
    normal_font = Font(name=font_name, size=11)
    title_font = Font(name=font_name, size=14, bold=True)
    note_font = Font(name=font_name, size=9, italic=True, color="666666")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    num_cols = 8

    ws["A1"] = f"Southwest Airlines — One-Way Flight Deals from {origin_label} ({origin_code})"
    ws["A1"].font = title_font
    ws.merge_cells("A1:H1")

    ws["A2"] = (
        f"Live data fetched from southwest.com — captured {datetime.now().strftime('%m/%d/%Y %H:%M')}. "
        f"Distances are great-circle (straight-line) miles between airports, not flown route miles."
    )
    ws["A2"].font = note_font
    ws.merge_cells("A2:H2")

    headers = ["City", "State", "Airport", f"Distance from {origin_code} (mi)", "Departing", "One-Way Price ($)", "Miles per $", "Rank (Cheapest = 1)"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    start_row = header_row + 1
    for i, (_rank, city, state, code, dist, date_obj, price, _mpd) in enumerate(rows):
        r = start_row + i
        ws.cell(row=r, column=1, value=city).font = normal_font
        ws.cell(row=r, column=2, value=state).font = normal_font
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=code).font = normal_font
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
        dist_cell = ws.cell(row=r, column=4, value=dist)
        dist_cell.font = normal_font
        dist_cell.number_format = "#,##0"
        dist_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=fmt_date(date_obj)).font = normal_font
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        price_cell = ws.cell(row=r, column=6, value=price)
        price_cell.font = normal_font
        price_cell.number_format = "$#,##0"
        price_cell.alignment = Alignment(horizontal="center")
        # Live formulas, not the Python-computed values, so the sheet stays
        # correct if someone edits distance/price cells by hand in Excel.
        mpd_cell = ws.cell(row=r, column=7, value=f"=D{r}/F{r}")  # Distance / Price
        mpd_cell.font = normal_font
        mpd_cell.number_format = "0.00"
        mpd_cell.alignment = Alignment(horizontal="center")
        # RANK() reads the whole price column, so it's correct regardless of
        # which order the rows were written in (e.g. exported while sorted
        # by distance) — it doesn't depend on this row's position.
        rank_cell = ws.cell(row=r, column=8, value=f"=RANK(F{r},$F${start_row}:$F${start_row + len(rows) - 1},1)")
        rank_cell.font = normal_font
        rank_cell.alignment = Alignment(horizontal="center")
        for col in range(1, num_cols + 1):
            ws.cell(row=r, column=col).border = border
        if i % 2 == 1:
            for col in range(1, num_cols + 1):
                ws.cell(row=r, column=col).fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    end_row = start_row + len(rows) - 1

    summary_row = end_row + 2
    ws.cell(row=summary_row, column=1, value="Summary").font = Font(name=font_name, size=11, bold=True)
    labels = ["Number of destinations", "Cheapest fare ($)", "Most expensive fare ($)", "Average fare ($)",
              "Shortest distance (mi)", "Longest distance (mi)", "Average distance (mi)",
              "Average miles per $"]
    formulas = [
        f"=COUNTA(A{start_row}:A{end_row})",
        f"=MIN(F{start_row}:F{end_row})",
        f"=MAX(F{start_row}:F{end_row})",
        f"=AVERAGE(F{start_row}:F{end_row})",
        f"=MIN(D{start_row}:D{end_row})",
        f"=MAX(D{start_row}:D{end_row})",
        f"=AVERAGE(D{start_row}:D{end_row})",
        f"=AVERAGE(G{start_row}:G{end_row})",
    ]
    for i, (label, formula) in enumerate(zip(labels, formulas)):
        r = summary_row + 1 + i
        ws.cell(row=r, column=1, value=label).font = normal_font
        fc = ws.cell(row=r, column=2, value=formula)
        fc.font = Font(name=font_name, size=11, bold=True)
        if "fare" in label:
            fc.number_format = "$#,##0.00" if "Average" in label else "$#,##0"
        elif "distance" in label:
            fc.number_format = "#,##0"
        elif "miles per $" in label:
            fc.number_format = "0.00"

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 18
    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = f"A{start_row}"

    wb.save(path)


class PriceHistoryWindow(tk.Toplevel):
    """Step chart of one route's recorded prices over time, drawn on a plain Tk
    Canvas (no matplotlib, which would add ~30 MB to the .exe).

    Only price *changes* are stored, so a price holds until the next recorded
    point: hence a step line, extended with a dashed segment out to "now" to
    show the latest price is still the current one."""

    MARGIN = {"left": 70, "right": 30, "top": 20, "bottom": 50}
    LINE_COLOR = "#004C97"

    def __init__(self, master, title, history):
        super().__init__(master)
        self.title(title)
        self.geometry("760x500")
        self.minsize(480, 340)
        self.history = history  # [(recorded_at, price, depart_date_iso), ...] oldest first
        self.hover_points = []  # (x, y, tooltip text) for each recorded point

        if history:
            prices = [price for _at, price, _dep in history]
            summary = (
                f"Current ${history[-1][1]:,}    Lowest ${min(prices):,}    Highest ${max(prices):,}    "
                f"{len(history)} recorded price(s) since {history[0][0]:%m/%d/%Y}"
            )
        else:
            summary = "No price history recorded yet for this route."
        ttk.Label(self, text=summary, padding=(10, 8)).pack(fill="x")

        self.canvas = tk.Canvas(self, background="white", highlightthickness=0)
        self.canvas.pack(fill="both", expand=True, padx=10)
        ttk.Label(
            self,
            text="Prices are recorded only when they change. The dashed line shows the latest price is still current.",
            foreground="#666666", padding=(10, 6),
        ).pack(fill="x")

        self.canvas.bind("<Configure>", lambda _event: self.redraw())
        self.canvas.bind("<Motion>", self.on_motion)
        self.canvas.bind("<Leave>", lambda _event: self.canvas.delete("tooltip"))

    @staticmethod
    def _nice_axis(lo, hi):
        """Round (lo, hi) outward to tidy dollar tick marks; returns (axis_min, axis_max, step)."""
        pad = max((hi - lo) * 0.1, 5)  # breathing room, and a visible axis even for a flat price
        lo, hi = max(0, lo - pad), hi + pad
        raw_step = (hi - lo) / 5
        magnitude = 10 ** math.floor(math.log10(raw_step))
        step = next(factor * magnitude for factor in (1, 2, 5, 10) if factor * magnitude >= raw_step)
        return math.floor(lo / step) * step, math.ceil(hi / step) * step, step

    def redraw(self):
        c = self.canvas
        c.delete("all")
        self.hover_points = []
        width, height = c.winfo_width(), c.winfo_height()
        left, right = self.MARGIN["left"], width - self.MARGIN["right"]
        top, bottom = self.MARGIN["top"], height - self.MARGIN["bottom"]

        if not self.history:
            c.create_text(width / 2, height / 2, justify="center", fill="#666666",
                          text="No price history recorded yet.\nIt builds up as you fetch.")
            return
        if right - left < 50 or bottom - top < 50:
            return  # window too small to draw anything useful

        t_start = self.history[0][0]
        t_now = max(datetime.now(), self.history[-1][0])
        # Show at least an hour so a brand-new route (one point, recorded minutes ago)
        # gets distinct axis labels; the dashed "still current" line still stops at now.
        t_end = max(t_now, t_start + timedelta(hours=1))
        span = (t_end - t_start).total_seconds()

        prices = [price for _at, price, _dep in self.history]
        y_min, y_max, step = self._nice_axis(min(prices), max(prices))

        inset = 12  # keeps the first/last dots and their labels off the axis lines

        def x_of(when):
            return left + inset + (right - left - 2 * inset) * (when - t_start).total_seconds() / span

        def y_of(price):
            return bottom - (bottom - top) * (price - y_min) / (y_max - y_min)

        value = y_min
        while value <= y_max + 1e-9:
            y = y_of(value)
            c.create_line(left, y, right, y, fill="#e6e6e6")
            c.create_text(left - 8, y, text=f"${value:,.0f}", anchor="e", fill="#444444")
            value += step

        date_format = "%m/%d\n%H:%M" if span < 2 * 86400 else "%m/%d/%y"
        for i in range(5):
            when = t_start + (t_end - t_start) * i / 4
            x = x_of(when)
            c.create_line(x, bottom, x, bottom + 5, fill="#888888")
            c.create_text(x, bottom + 8, text=when.strftime(date_format), anchor="n",
                          fill="#444444", justify="center")

        c.create_line(left, top, left, bottom, fill="#888888")
        c.create_line(left, bottom, right, bottom, fill="#888888")

        # Step path: hold the old price until each change, then jump to the new one.
        coords = []
        held_y = None
        for when, price, _dep in self.history:
            x, y = x_of(when), y_of(price)
            if held_y is not None:
                coords += [x, held_y]
            coords += [x, y]
            held_y = y
        if len(coords) >= 4:
            c.create_line(*coords, fill=self.LINE_COLOR, width=2)
        last_x, end_x = coords[-2], x_of(t_now)
        if end_x - last_x > 2:
            c.create_line(last_x, held_y, end_x, held_y, fill=self.LINE_COLOR, width=2, dash=(4, 3))

        label_dots = len(self.history) <= 12  # beyond that, labels would pile up; use hover instead
        for when, price, depart in self.history:
            x, y = x_of(when), y_of(price)
            c.create_oval(x - 4, y - 4, x + 4, y + 4, fill=self.LINE_COLOR, outline="white")
            if label_dots:
                c.create_text(x, y - 10, text=f"${price:,}", anchor="s", fill=self.LINE_COLOR)
            tip = f"${price:,}\n{when:%b %d, %Y %H:%M}"
            if depart:
                tip += f"\nDeparting {fmt_date(datetime.fromisoformat(depart))}"
            self.hover_points.append((x, y, tip))

    def on_motion(self, event):
        c = self.canvas
        c.delete("tooltip")
        if not self.hover_points:
            return
        x, y, tip = min(self.hover_points, key=lambda p: (p[0] - event.x) ** 2 + (p[1] - event.y) ** 2)
        if (x - event.x) ** 2 + (y - event.y) ** 2 > 14 ** 2:
            return
        text = c.create_text(x + 12, y - 12, text=tip, anchor="sw", justify="left", tags="tooltip")
        x1, y1, x2, y2 = c.bbox(text)
        c.move(text, min(0, c.winfo_width() - 6 - x2), max(0, 6 - y1))  # keep it inside the canvas
        x1, y1, x2, y2 = c.bbox(text)
        box = c.create_rectangle(x1 - 4, y1 - 3, x2 + 4, y2 + 3, fill="#ffffe0", outline="#999999", tags="tooltip")
        c.tag_lower(box, text)


class WatchDialog(tk.Toplevel):
    """Set (or edit) the price target for one route. `on_save(target)` is called
    with a validated number; `on_remove()` (optional) removes an existing watch."""

    def __init__(self, master, origin, dest, city, state, current_price, existing_target, on_save, on_remove=None):
        super().__init__(master)
        self.title("Edit Price Watch" if existing_target is not None else "Set Price Watch")
        self.resizable(False, False)
        self.transient(master)
        self.on_save, self.on_remove = on_save, on_remove

        body = ttk.Frame(self, padding=16)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text=f"{origin} → {dest}   {city}, {state}", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=3, sticky="w")
        current = fmt_money(current_price) if current_price is not None else "unknown (fetch first)"
        ttk.Label(body, text=f"Current price: {current}").grid(row=1, column=0, columnspan=3, sticky="w", pady=(4, 12))

        ttk.Label(body, text="Notify me when the price is at or below  $").grid(row=2, column=0, sticky="w")
        self.target_var = tk.StringVar(value="" if existing_target is None else f"{existing_target:g}")
        self.entry = ttk.Entry(body, textvariable=self.target_var, width=8)
        self.entry.grid(row=2, column=1, sticky="w", padx=(2, 0))
        self.error_var = tk.StringVar()
        ttk.Label(body, textvariable=self.error_var, foreground="#b00020").grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(6, 0))

        buttons = ttk.Frame(body)
        buttons.grid(row=4, column=0, columnspan=3, sticky="e", pady=(14, 0))
        if on_remove is not None:
            ttk.Button(buttons, text="Remove Watch", command=self._remove).pack(side="left", padx=(0, 16))
        ttk.Button(buttons, text="Save", command=self.save).pack(side="left")
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="left", padx=(8, 0))

        self.entry.focus_set()
        self.bind("<Return>", lambda _e: self.save())
        self.bind("<Escape>", lambda _e: self.destroy())

    def save(self):
        text = self.target_var.get().strip().lstrip("$").replace(",", "")
        try:
            target = float(text)
        except ValueError:
            self.error_var.set("Enter a dollar amount, e.g. 90")
            return
        if not 0 < target <= 100000:
            self.error_var.set("Enter an amount greater than $0.")
            return
        self.destroy()
        self.on_save(int(target) if target.is_integer() else target)

    def _remove(self):
        self.destroy()
        self.on_remove()


class ChatPickerDialog(tk.Toplevel):
    """Choose which detected Telegram chat (private chat, group, or channel) to send alerts to."""

    def __init__(self, master, chats, on_pick):
        super().__init__(master)
        self.title("Choose a Chat")
        self.resizable(False, False)
        self.transient(master)
        self.chats, self.on_pick = chats, on_pick

        body = ttk.Frame(self, padding=16)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Telegram found these chats. Which should get the alerts?").pack(anchor="w")
        self.listbox = tk.Listbox(body, width=62, height=min(max(len(chats), 3), 10), activestyle="none", exportselection=False)
        for chat_id, name, kind in chats:
            self.listbox.insert("end", f"{name}   ({kind})   {chat_id}")
        self.listbox.selection_set(0)
        self.listbox.pack(pady=(8, 12))
        self.listbox.bind("<Double-1>", lambda _e: self.choose())

        buttons = ttk.Frame(body)
        buttons.pack(anchor="e")
        ttk.Button(buttons, text="Use Selected", command=self.choose).pack(side="left")
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="left", padx=(8, 0))
        self.bind("<Return>", lambda _e: self.choose())
        self.bind("<Escape>", lambda _e: self.destroy())

    def choose(self):
        selection = self.listbox.curselection()
        if selection:
            chat = self.chats[selection[0]]
            self.destroy()
            self.on_pick(chat)


class TelegramSetupWindow(tk.Toplevel):
    """Configure the Telegram bot (token + chat ID), test it, and save to settings.json."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Telegram Setup")
        self.resizable(False, False)
        self.transient(app)

        body = ttk.Frame(self, padding=16)
        body.pack(fill="both", expand=True)
        steps = (
            "How to set up alerts:\n"
            "1. In Telegram, message @BotFather, send /newbot, and follow the prompts.\n"
            "    It gives you a bot token — paste it below.\n"
            "2. Send your new bot any message (e.g. \"hi\").\n"
            "    To alert a channel or group instead (so others get them too), add the bot to it —\n"
            "    for a channel, as an administrator who can post messages — and post a message there.\n"
            "3. Click Detect Chat ID and pick the chat (or type a chat ID in yourself).\n"
            "4. Click Send Test Message, then Save."
        )
        ttk.Label(body, text=steps, justify="left").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 12))

        saved = app.settings["telegram"]
        ttk.Label(body, text="Bot token:").grid(row=1, column=0, sticky="w")
        self.token_var = tk.StringVar(value=saved["bot_token"])
        self.token_entry = ttk.Entry(body, textvariable=self.token_var, width=48, show="•")
        self.token_entry.grid(row=1, column=1, sticky="w", padx=(8, 0))
        self.show_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(body, text="Show", variable=self.show_var, command=self._toggle_show).grid(
            row=1, column=2, sticky="w", padx=(8, 0))

        ttk.Label(body, text="Chat ID:").grid(row=2, column=0, sticky="w", pady=(8, 0))
        self.chat_var = tk.StringVar(value=saved["chat_id"])
        ttk.Entry(body, textvariable=self.chat_var, width=24).grid(row=2, column=1, sticky="w", padx=(8, 0), pady=(8, 0))
        ttk.Button(body, text="Detect Chat ID", command=self.detect).grid(row=2, column=2, padx=(8, 0), pady=(8, 0))

        self.result_var = tk.StringVar()
        self.result_label = ttk.Label(body, textvariable=self.result_var, wraplength=520, justify="left")
        self.result_label.grid(row=3, column=0, columnspan=3, sticky="w", pady=(12, 0))
        ttk.Label(
            body, foreground="#666666", wraplength=520, justify="left",
            text="The token is stored in plain text in settings.json next to the app. Keep that file private "
                 "(it's excluded from git), and don't share it.",
        ).grid(row=4, column=0, columnspan=3, sticky="w", pady=(8, 0))

        buttons = ttk.Frame(body)
        buttons.grid(row=5, column=0, columnspan=3, sticky="e", pady=(14, 0))
        ttk.Button(buttons, text="Send Test Message", command=self.send_test).pack(side="left")
        ttk.Button(buttons, text="Save", command=self.save).pack(side="left", padx=(16, 0))
        ttk.Button(buttons, text="Close", command=self.destroy).pack(side="left", padx=(8, 0))

    def _toggle_show(self):
        self.token_entry.config(show="" if self.show_var.get() else "•")

    def _say(self, text, ok=True):
        self.result_var.set(text)
        self.result_label.config(foreground="#1b7f3b" if ok else "#b00020")
        self.update_idletasks()

    def _credentials(self, need_chat=True):
        token, chat = self.token_var.get().strip(), self.chat_var.get().strip()
        if not token:
            self._say("Enter the bot token first.", ok=False)
            return None
        if need_chat and not chat:
            self._say("Enter a chat ID, or click Detect Chat ID.", ok=False)
            return None
        return token, chat

    def detect(self):
        creds = self._credentials(need_chat=False)
        if not creds:
            return
        self._say("Contacting Telegram...")
        try:
            chats = detect_telegram_chats(creds[0])
        except TelegramError as exc:
            self._say(str(exc), ok=False)
            return
        if len(chats) == 1:
            self._use_chat(chats[0])
        else:
            self._say(f"Found {len(chats)} chats - choose one.")
            ChatPickerDialog(self, chats, self._use_chat)

    def _use_chat(self, chat):
        chat_id, name, kind = chat
        self.chat_var.set(chat_id)
        self._say(f"Using {name} ({kind}, {chat_id}). Click Send Test Message to confirm, then Save.")

    def send_test(self):
        creds = self._credentials()
        if not creds:
            return
        self._say("Sending...")
        try:
            send_telegram(creds[0], creds[1], f"✅ {APP_NAME}: Telegram alerts are set up.")
        except TelegramError as exc:
            self._say(f"Test failed: {exc}", ok=False)
            return
        self._say("Test message sent — check Telegram. Click Save to keep these settings.")

    def save(self):
        try:
            self.app.save_telegram_settings(self.token_var.get().strip(), self.chat_var.get().strip())
        except OSError as exc:
            self._say(f"Couldn't save settings.json: {exc}", ok=False)
            return
        self._say("Saved to settings.json.")


class WatchManagerWindow(tk.Toplevel):
    """Lists every price watch with its current price and status; edit or remove them."""

    COLUMNS = (("route", "Route", 90), ("dest", "Destination", 190), ("target", "Alert at or below", 120),
               ("current", "Current", 80), ("status", "Status", 420))

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Price Watches")
        self.geometry("960x360")
        self.minsize(700, 260)

        self.telegram_var = tk.StringVar()
        ttk.Label(self, textvariable=self.telegram_var, padding=(10, 8)).pack(fill="x")

        frame = ttk.Frame(self)
        frame.pack(fill="both", expand=True, padx=10)
        self.tree = ttk.Treeview(frame, columns=[c[0] for c in self.COLUMNS], show="headings", selectmode="browse")
        for key, label, width in self.COLUMNS:
            self.tree.heading(key, text=label)
            self.tree.column(key, width=width, anchor="w" if key in ("dest", "status") else "center")
        vsb = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self.tree.bind("<Double-1>", lambda _e: self.edit())

        self.empty_var = tk.StringVar()
        ttk.Label(self, textvariable=self.empty_var, foreground="#666666", padding=(10, 6)).pack(fill="x")

        buttons = ttk.Frame(self, padding=(10, 0, 10, 10))
        buttons.pack(fill="x")
        ttk.Button(buttons, text="Edit Target...", command=self.edit).pack(side="left")
        ttk.Button(buttons, text="Remove", command=self.remove).pack(side="left", padx=(8, 0))
        ttk.Button(buttons, text="Telegram Setup...", command=app.open_telegram_setup).pack(side="left", padx=(24, 0))
        ttk.Button(buttons, text="Close", command=self.destroy).pack(side="right")
        self.refresh()

    def refresh(self):
        selected = self.tree.selection()
        self.tree.delete(*self.tree.get_children())
        for index, watch in enumerate(self.app.settings["watches"]):
            current, status = self.app.describe_watch(watch)
            self.tree.insert("", "end", iid=str(index), values=(
                f"{watch['origin']} → {watch['dest']}", f"{watch['city']}, {watch['state']}".strip(", "),
                f"≤ {fmt_money(watch['target'])}", fmt_money(current) if current is not None else "—", status))
        if selected and self.tree.exists(selected[0]):
            self.tree.selection_set(selected[0])
        cfg = self.app.settings["telegram"]
        self.telegram_var.set(
            "Telegram: set up — alerts will be sent." if cfg["bot_token"] and cfg["chat_id"]
            else "Telegram: NOT set up — alerts can't be sent until you configure it (button below)."
        )
        self.empty_var.set(
            "" if self.app.settings["watches"]
            else "No price watches yet. Select a row in the main table (or right-click it) and choose Set Price Watch."
        )

    def _selected_watch(self):
        selection = self.tree.selection()
        if not selection:
            messagebox.showinfo("Price Watches", "Select a watch first.", parent=self)
            return None
        return self.app.settings["watches"][int(selection[0])]

    def edit(self):
        watch = self._selected_watch()
        if watch:
            self.app.open_watch_dialog(watch["origin"], watch["dest"], watch["city"], watch["state"])

    def remove(self):
        watch = self._selected_watch()
        if watch and messagebox.askyesno(
                "Remove watch", f"Stop watching {watch['origin']} → {watch['dest']}?", parent=self):
            self.app.remove_watch(watch["origin"], watch["dest"])


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_NAME} v{APP_VERSION}")
        # Wide enough for all 9 result columns (their widths sum to ~960px)
        # plus the scrollbar and window padding.
        self.geometry("1100x560")

        menubar = tk.Menu(self)
        alerts_menu = tk.Menu(menubar, tearoff=0)
        alerts_menu.add_command(label="Set Price Watch for Selected Row...", command=self.set_watch_for_selected_row)
        alerts_menu.add_command(label="View Price Watches...", command=self.open_watch_manager)
        alerts_menu.add_separator()
        alerts_menu.add_command(label="Telegram Setup...", command=self.open_telegram_setup)
        menubar.add_cascade(label="Alerts", menu=alerts_menu)
        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="About", command=self.show_about)
        menubar.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menubar)

        # Price watches + Telegram config, persisted in settings.json.
        self.settings, settings_warning = load_settings()
        self.pending_alerts = set()  # (origin, dest) keys with a Telegram send in flight
        self.watch_errors = {}       # (origin, dest) -> last send-failure text, until it succeeds
        self.notify_queue = queue.Queue()
        self.notify_outstanding = 0
        self.notify_polling = False
        self.alert_note = ("", False)  # (text, is_error) shown in the alert label until replaced
        self.watch_manager = None
        self.telegram_setup = None

        self.rows = []
        self.display_rows = []
        self.sort_column = None
        self.sort_reverse = False
        self.price_changes = {}  # airport code -> "up" / "down", vs. the previous fetch
        self.price_deltas = {}   # airport code -> dollar change vs. the previous fetch
        self.result_queue = queue.Queue()
        self.fetching = False
        self.auto_refresh_active = False
        self.auto_refresh_job = None  # after() id for the next scheduled tick, so it can be cancelled
        self.auto_refresh_ms = None

        self.origin_code = DEFAULT_ORIGIN
        self.origin_label = ""
        self.previous_origin_code = None  # last-fetched origin, so price-change diffs don't cross origins
        self.origin_lookup = {}  # dropdown display string -> airport code
        self.origins_queue = queue.Queue()
        self.fetch_is_auto = False  # True when the running fetch was started by auto-refresh
        self.network_retry_job = None  # after() id of a pending quiet retry (fares)
        self.origins_retry_job = None  # after() id of a pending retry (airport list)
        self.last_pdf_path = None
        self.last_xlsx_path = None

        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        origin_frame = ttk.Frame(self, padding=(10, 0))
        origin_frame.pack(fill="x")
        ttk.Label(origin_frame, text="Home airport:").pack(side="left")
        # Bootstrapped with just BWI so the dropdown is usable immediately;
        # start_load_origins() fills in the full list in the background.
        bootstrap_label = f"{DEFAULT_ORIGIN} — Baltimore/Washington, MD"
        self.origin_lookup[bootstrap_label] = DEFAULT_ORIGIN
        self.origin_var = tk.StringVar(value=bootstrap_label)
        self.origin_combo = ttk.Combobox(
            origin_frame, textvariable=self.origin_var, state="readonly",
            width=40, values=[bootstrap_label],
        )
        self.origin_combo.pack(side="left", padx=(4, 0))
        self.origin_combo.bind("<<ComboboxSelected>>", self.on_origin_selected)
        # Short alert status on the right of this row (watch count / last alert sent / failure).
        self.alert_label = ttk.Label(origin_frame, text="", foreground="#444444")
        self.alert_label.pack(side="right")

        self.fetch_btn = ttk.Button(top, text="Fetch Latest Deals", command=self.start_fetch)
        self.fetch_btn.pack(side="left")

        self.export_pdf_btn = ttk.Button(top, text="Export to PDF", command=self.do_export_pdf, state="disabled")
        self.export_pdf_btn.pack(side="left", padx=(8, 0))

        self.open_pdf_btn = ttk.Button(top, text="Open PDF", command=self.open_pdf, state="disabled")
        self.open_pdf_btn.pack(side="left", padx=(8, 0))

        self.export_xlsx_btn = ttk.Button(top, text="Export to XLSX", command=self.do_export_xlsx, state="disabled")
        self.export_xlsx_btn.pack(side="left", padx=(8, 0))

        self.open_xlsx_btn = ttk.Button(top, text="Open XLSX", command=self.open_xlsx, state="disabled")
        self.open_xlsx_btn.pack(side="left", padx=(8, 0))

        self.progress = ttk.Progressbar(top, mode="indeterminate", length=150)
        self.progress.pack(side="right")

        refresh = ttk.Frame(self, padding=(10, 0))
        refresh.pack(fill="x")
        ttk.Label(refresh, text="Auto-refresh every:").pack(side="left")
        # Blank by default — auto-refresh stays off until a value is entered
        # and "Start Auto-Refresh" is clicked.
        self.interval_var = tk.StringVar(value="")
        self.interval_entry = ttk.Entry(refresh, textvariable=self.interval_var, width=6)
        self.interval_entry.pack(side="left", padx=(4, 4))
        ttk.Label(refresh, text="minutes").pack(side="left")
        self.auto_refresh_btn = ttk.Button(refresh, text="Start Auto-Refresh", command=self.toggle_auto_refresh)
        self.auto_refresh_btn.pack(side="left", padx=(8, 0))
        ttk.Label(refresh, text="Tip: double-click a row for its price history  |  right-click to set a price watch",
                  foreground="#666666").pack(side="right")

        self.status_var = tk.StringVar(value="Click \"Fetch Latest Deals\" to load current fares from southwest.com.")
        # Packed with side="bottom" *before* the tree so it reserves the
        # bottom strip of the window; the expanding tree then fills the rest.
        ttk.Label(self, textvariable=self.status_var, padding=(10, 4)).pack(side="bottom", fill="x")

        self.protocol("WM_DELETE_WINDOW", self._on_close)

        columns = ("rank", "city", "state", "airport", "distance", "departing", "price", "change", "mpd")
        self.tree = ttk.Treeview(self, columns=columns, show="headings")
        self.column_labels = {
            "rank": ("Rank", 50),
            "city": ("City", 190),
            "state": ("State", 60),
            "airport": ("Airport", 70),
            "distance": (f"Distance from {DEFAULT_ORIGIN} (mi)", 150),
            "departing": ("Departing", 90),
            "price": ("One-Way Price ($)", 120),
            "change": ("Change vs. Last Fetch", 140),
            "mpd": ("Miles per $", 90),
        }
        for col, (label, width) in self.column_labels.items():
            self.tree.heading(col, text=label, command=lambda c=col: self.sort_by(c))
            anchor = "w" if col == "city" else "center"
            self.tree.column(col, width=width, anchor=anchor)

        self.tree.tag_configure("price_up", background="#f8d7da")    # light red
        self.tree.tag_configure("price_down", background="#d4edda")  # light green
        self.tree.bind("<Double-1>", self.on_row_double_click)
        self.row_menu = tk.Menu(self, tearoff=0)
        self.row_menu.add_command(label="Set Price Watch...", command=self.set_watch_for_selected_row)
        self.row_menu.add_command(label="View Price History", command=self.show_history_for_selected_row)
        self.tree.bind("<Button-3>", self.on_row_right_click)

        vsb = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True, padx=(10, 0), pady=10)
        vsb.pack(side="right", fill="y", pady=10)

        self._update_alert_label()
        if settings_warning:
            # Rare (unreadable settings.json); a popup so it isn't lost in the status bar.
            self.after(300, lambda: messagebox.showwarning("Settings", settings_warning))
        self.start_load_origins()

    def start_load_origins(self):
        threading.Thread(target=self._load_origins_worker, daemon=True).start()
        self.after(100, self._poll_origins_queue)

    def _load_origins_worker(self):
        try:
            airports = fetch_origin_airports()
            self.origins_queue.put(("ok", airports))
        except Exception as exc:
            kind = "network_error" if is_network_error(exc) else "error"
            self.origins_queue.put((kind, str(exc)))

    def _poll_origins_queue(self):
        try:
            status, payload = self.origins_queue.get_nowait()
        except queue.Empty:
            self.after(200, self._poll_origins_queue)
            return

        if status == "network_error":
            # Typically the network isn't back yet after standby/login — keep
            # trying quietly so the dropdown fills in without a restart.
            self.status_var.set(
                f"Couldn't load the full airport list ({payload}); showing BWI only. "
                f"Retrying in {NETWORK_RETRY_SECONDS}s..."
            )
            self.origins_retry_job = self.after(NETWORK_RETRY_SECONDS * 1000, self.start_load_origins)
            return
        if status == "error":
            # Non-fatal — the BWI-only bootstrap entry still works fine.
            self.status_var.set(f"Couldn't load the full airport list ({payload}); showing BWI only.")
            return

        lookup = {}
        for code, name in payload:
            # Southwest's display names are "City, ST - CODE"; drop the
            # " - CODE" suffix since the code is already shown up front.
            city_state = name.rsplit(" - ", 1)[0] if " - " in name else name
            lookup[f"{code} — {city_state}"] = code

        self.origin_lookup = lookup
        labels = sorted(lookup)
        self.origin_combo.config(values=labels)

        # Keep whatever's currently selected (still just BWI at this point,
        # unless the user was fast) pointed at the right dropdown entry.
        current = next((label for label, code in lookup.items() if code == self.origin_code), None)
        if current:
            self.origin_var.set(current)

        self.status_var.set(f"Loaded {len(labels)} home airports. Click \"Fetch Latest Deals\" to load current fares from southwest.com.")

    def on_origin_selected(self, _event=None):
        label = self.origin_var.get()
        self.origin_code = self.origin_lookup.get(label, label.split("—", 1)[0].strip())

    def start_fetch(self, auto=False):
        self._cancel_network_retry()  # a fetch is starting now, so any pending retry is redundant
        self.fetching = True
        self.fetch_is_auto = auto
        self.fetch_btn.config(state="disabled")
        self.export_pdf_btn.config(state="disabled")
        self.open_pdf_btn.config(state="disabled")
        self.export_xlsx_btn.config(state="disabled")
        self.open_xlsx_btn.config(state="disabled")
        self.status_var.set("Fetching live fares from southwest.com...")
        self.progress.start(12)
        threading.Thread(target=self._fetch_worker, daemon=True).start()
        self.after(100, self._poll_queue)

    def _fetch_worker(self):
        # Runs on a background thread — network I/O here so the Tk mainloop
        # (and the window) stays responsive. Tkinter widgets aren't
        # thread-safe, so this thread must never touch self.tree/self.rows
        # etc. directly; it only hands the result to _poll_queue via the
        # queue, which runs back on the main thread.
        origin_code = self.origin_code  # snapshot: whatever's selected right now
        try:
            origin_label, rows = fetch_deals(origin_code)
            self.result_queue.put(("ok", origin_code, origin_label, rows))
        except Exception as exc:
            kind = "network_error" if is_network_error(exc) else "error"
            self.result_queue.put((kind, str(exc)))

    def _cancel_network_retry(self):
        if self.network_retry_job is not None:
            self.after_cancel(self.network_retry_job)
            self.network_retry_job = None

    def _network_retry(self):
        self.network_retry_job = None
        if self.auto_refresh_active and not self.fetching:
            self.start_fetch(auto=True)

    def _poll_queue(self):
        # Polls (rather than blocking) so the mainloop keeps servicing the
        # UI while we wait; reschedules itself every 100ms until the
        # worker thread above has put something in the queue.
        try:
            result = self.result_queue.get_nowait()
        except queue.Empty:
            self.after(100, self._poll_queue)
            return

        self.progress.stop()
        self.fetch_btn.config(state="normal")
        self.fetching = False

        if result[0] == "network_error":
            detail = result[1]
            if self.fetch_is_auto and self.auto_refresh_active:
                # Unattended refresh (often firing the instant the PC wakes from
                # standby, before Wi-Fi/DNS is back): no popup, just say so and
                # try again shortly. The table keeps showing the last good data.
                self.status_var.set(
                    f"Couldn't reach southwest.com ({detail}) — network may still be "
                    f"reconnecting. Retrying in {NETWORK_RETRY_SECONDS}s..."
                )
                self.network_retry_job = self.after(NETWORK_RETRY_SECONDS * 1000, self._network_retry)
                return
            self.status_var.set(f"Couldn't reach southwest.com ({detail})")
            messagebox.showerror(
                "Fetch failed",
                "Couldn't reach southwest.com. Check your internet connection — it can "
                "take a little while to come back after standby or login — then try again."
                f"\n\n{detail}",
            )
            return
        if result[0] == "error":
            message = result[1]
            self.status_var.set(f"Error: {message}")
            messagebox.showerror("Fetch failed", message)
            return

        _, origin_code, origin_label, payload = result

        # Snapshot the OLD self.rows before we overwrite it below — this is
        # the "previous fetch" we compare the new prices against. On the
        # very first fetch (or after switching to a different home airport)
        # there's nothing comparable, so previous_prices stays empty and
        # nothing gets highlighted.
        previous_prices = {}
        if self.previous_origin_code == origin_code:
            previous_prices = {row[3]: row[6] for row in self.rows}  # airport code -> price

        self.origin_code = origin_code
        self.origin_label = origin_label
        self.previous_origin_code = origin_code
        self.column_labels["distance"] = (f"Distance from {origin_code} (mi)", self.column_labels["distance"][1])
        self.rows = payload
        self.display_rows = list(payload)
        self.sort_column = None
        self.sort_reverse = False

        # Diff by airport code (not by row position — sort order/rank can
        # differ between fetches) and only flag actual increases/decreases.
        # Codes with no prior price (new to this fetch) are left untouched
        # in price_changes, which populate_tree treats as "no highlight".
        self.price_changes = {}
        self.price_deltas = {}  # airport code -> new price minus previous price (only codes seen last fetch)
        up = down = same = 0
        for row in self.rows:
            code, price = row[3], row[6]
            if code not in previous_prices:
                continue
            self.price_deltas[code] = price - previous_prices[code]
            if price > previous_prices[code]:
                self.price_changes[code] = "up"
                up += 1
            elif price < previous_prices[code]:
                self.price_changes[code] = "down"
                down += 1
            else:
                same += 1

        for col, (label, _width) in self.column_labels.items():
            self.tree.heading(col, text=label)
        self.populate_tree()
        self.export_pdf_btn.config(state="normal")
        self.export_xlsx_btn.config(state="normal")

        timestamp = datetime.now().strftime("%m/%d/%Y %H:%M")
        status = f"Loaded {len(self.rows)} destinations from {origin_label} ({origin_code}) as of {timestamp}."
        if previous_prices:
            status += f" Since last fetch: {up} up, {down} down, {same} unchanged."

        # Persist to the local price history (only prices that changed since the
        # last stored one). A failure here shouldn't fail the fetch itself.
        try:
            saved = record_prices(origin_code, payload)
            if saved:
                status += f" Saved {saved} price(s) to history."
        except (sqlite3.Error, OSError) as exc:
            status += f" (Couldn't save price history: {exc})"
        self.status_var.set(status)

        self._evaluate_watches(origin_code, payload)

    def on_row_double_click(self, event):
        if self.tree.identify_region(event.x, event.y) != "cell":
            return  # a double-click on a column header or empty space
        item = self.tree.identify_row(event.y)
        if item:
            self._show_history_for_item(item)

    def on_row_right_click(self, event):
        item = self.tree.identify_row(event.y)
        if not item or self.tree.identify_region(event.x, event.y) != "cell":
            return
        self.tree.selection_set(item)
        try:
            self.row_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.row_menu.grab_release()

    def _show_history_for_item(self, item):
        values = self.tree.item(item, "values")
        city, state, dest = values[1], values[2], values[3]
        # previous_origin_code is the origin of the rows currently in the table;
        # self.origin_code may already point at a newer dropdown selection.
        origin = self.previous_origin_code
        try:
            history = load_price_history(origin, dest)
        except (sqlite3.Error, OSError) as exc:
            messagebox.showerror("Price history", f"Couldn't read the price history: {exc}")
            return
        PriceHistoryWindow(self, f"Price history — {city}, {state} ({origin} → {dest})", history)

    def _selected_route(self):
        """(origin, dest, city, state) for the selected table row, or None (after telling the user)."""
        selection = self.tree.selection()
        if not selection or self.previous_origin_code is None:
            messagebox.showinfo("Price watch", "Fetch the fares first, then select a row.")
            return None
        values = self.tree.item(selection[0], "values")
        return self.previous_origin_code, values[3], values[1], values[2]

    def show_history_for_selected_row(self):
        if self._selected_route():
            self._show_history_for_item(self.tree.selection()[0])

    # --- Price watches & Telegram alerts -------------------------------------

    def watch_for(self, origin, dest):
        return next((w for w in self.settings["watches"] if w["origin"] == origin and w["dest"] == dest), None)

    def telegram_configured(self):
        cfg = self.settings["telegram"]
        return bool(cfg["bot_token"] and cfg["chat_id"])

    def current_price(self, origin, dest):
        """Latest known price: the loaded table if it's for this origin, else the
        newest entry in the price history (which only changes when the price does)."""
        if self.previous_origin_code == origin:
            for row in self.rows:
                if row[3] == dest:
                    return row[6]
        try:
            history = load_price_history(origin, dest)
        except (sqlite3.Error, OSError):
            return None
        return history[-1][1] if history else None

    def describe_watch(self, watch):
        """(current_price_or_None, status_text) for the Price Watches window."""
        key = (watch["origin"], watch["dest"])
        current = self.current_price(*key)
        if key in self.watch_errors:
            return current, f"Alert failed: {self.watch_errors[key]} (retries on next fetch)"
        if key in self.pending_alerts:
            return current, "Sending alert..."
        if current is None:
            return current, "Waiting for a price (fetch first)"
        if current <= watch["target"]:
            if watch.get("last_alert_price") is not None:
                try:
                    when = datetime.strptime(watch["last_alert_at"], HISTORY_TIME_FORMAT).strftime("%m/%d %H:%M")
                except (TypeError, ValueError):
                    when = "earlier"
                return current, f"At/below target - alert sent {when} at {fmt_money(watch['last_alert_price'])}"
            return current, "At/below target - alert pending (is Telegram set up?)"
        return current, f"Watching - {fmt_money(current - watch['target'])} above target"

    def _persist_settings(self):
        try:
            save_settings(self.settings)
            return True
        except OSError as exc:
            self._set_alert_note("Couldn't save settings.json", error=True)
            messagebox.showerror("Settings", f"Couldn't save settings.json: {exc}")
            return False

    def _set_alert_note(self, text, error=False):
        self.alert_note = (text, error)
        self._update_alert_label()

    def _update_alert_label(self):
        text, is_error = self.alert_note
        if not text:
            count = len(self.settings["watches"])
            if count:
                text = f"{count} price watch{'es' if count != 1 else ''} active"
                if not self.telegram_configured():
                    text += " (Telegram not set up)"
        self.alert_label.config(text=text, foreground="#b00020" if is_error else "#444444")

    def _after_watches_changed(self):
        if not self.settings["watches"] and not self.alert_note[1]:
            self.alert_note = ("", False)
        self._update_alert_label()
        if self.watch_manager is not None and self.watch_manager.winfo_exists():
            self.watch_manager.refresh()

    def open_watch_manager(self):
        if self.watch_manager is not None and self.watch_manager.winfo_exists():
            self.watch_manager.refresh()
            self.watch_manager.lift()
        else:
            self.watch_manager = WatchManagerWindow(self)

    def open_telegram_setup(self):
        if self.telegram_setup is not None and self.telegram_setup.winfo_exists():
            self.telegram_setup.lift()
        else:
            self.telegram_setup = TelegramSetupWindow(self)

    def save_telegram_settings(self, token, chat_id):
        self.settings["telegram"] = {"bot_token": token, "chat_id": chat_id}
        save_settings(self.settings)  # OSError propagates to the setup window, which reports it
        self.watch_errors.clear()  # old send failures may be fixed by the new settings
        self._after_watches_changed()
        if self.previous_origin_code is not None:
            # Watches that were already at/below target (waiting on Telegram) can alert now.
            self._evaluate_watches(self.previous_origin_code, self.rows)

    def set_watch_for_selected_row(self):
        route = self._selected_route()
        if route:
            self.open_watch_dialog(*route)

    def open_watch_dialog(self, origin, dest, city, state):
        existing = self.watch_for(origin, dest)
        WatchDialog(
            self, origin, dest, city, state, self.current_price(origin, dest),
            existing["target"] if existing else None,
            on_save=lambda target: self._save_watch(origin, dest, city, state, target),
            on_remove=(lambda: self.remove_watch(origin, dest)) if existing else None,
        )

    def _save_watch(self, origin, dest, city, state, target):
        watch = self.watch_for(origin, dest)
        if watch is None:
            watch = {"origin": origin, "dest": dest, "city": city, "state": state, "target": target,
                     "created": datetime.now().strftime(HISTORY_TIME_FORMAT),
                     "last_alert_price": None, "last_alert_at": None}
            self.settings["watches"].append(watch)
        else:
            watch["target"] = target
            watch["last_alert_price"] = watch["last_alert_at"] = None  # re-arm so the new target can alert right away
        self.watch_errors.pop((origin, dest), None)
        if not self._persist_settings():
            return
        self._after_watches_changed()
        if self.previous_origin_code == origin:
            self._evaluate_watches(origin, self.rows)  # the current price may already be at/below the target
        if not self.telegram_configured() and messagebox.askyesno(
                "Telegram not set up",
                "Watch saved, but alerts can't be sent until Telegram is set up. Set it up now?"):
            self.open_telegram_setup()

    def remove_watch(self, origin, dest):
        self.settings["watches"] = [
            w for w in self.settings["watches"] if not (w["origin"] == origin and w["dest"] == dest)
        ]
        self.watch_errors.pop((origin, dest), None)
        self.pending_alerts.discard((origin, dest))
        if self._persist_settings():
            self._after_watches_changed()

    def _evaluate_watches(self, origin_code, rows):
        """Compare each watch for this origin with the fresh prices and send alerts.

        Alert rules, so a watch doesn't spam you on every auto-refresh:
        - price at or below the target and no alert sent yet -> alert
        - already alerted -> alert again only if the price drops *below* the last alerted price
        - price back above the target -> re-arm (the next dip alerts again)
        A watch whose send failed (or Telegram isn't set up) isn't marked as alerted,
        so it's retried on the next fetch."""
        by_code = {row[3]: row for row in rows}
        jobs, unsent_hits, rearmed = [], [], False
        for watch in self.settings["watches"]:
            if watch["origin"] != origin_code or watch["dest"] not in by_code:
                continue
            row = by_code[watch["dest"]]
            key = (watch["origin"], watch["dest"])
            price, last = row[6], watch.get("last_alert_price")
            if price > watch["target"]:
                if last is not None:
                    watch["last_alert_price"] = watch["last_alert_at"] = None
                    rearmed = True
                continue
            if key in self.pending_alerts or (last is not None and price >= last):
                continue
            if self.telegram_configured():
                jobs.append((key, format_alert(watch, row), price))
            else:
                unsent_hits.append(f"{watch['dest']} {fmt_money(price)}")

        if rearmed:
            self._persist_settings()
        if jobs:
            self._start_notify(jobs)
        elif unsent_hits:
            self._set_alert_note(f"Watch hit ({', '.join(unsent_hits)}) - set up Telegram", error=True)
        self._after_watches_changed()

    def _start_notify(self, jobs):
        cfg = self.settings["telegram"]
        for key, _text, _price in jobs:
            self.pending_alerts.add(key)
        self.notify_outstanding += len(jobs)
        threading.Thread(
            target=self._notify_worker, args=(cfg["bot_token"], cfg["chat_id"], jobs), daemon=True
        ).start()
        if not self.notify_polling:
            self.notify_polling = True
            self.after(200, self._poll_notify_queue)

    def _notify_worker(self, token, chat_id, jobs):
        # Background thread: only talks to the network and the queue, never to widgets.
        for key, text, price in jobs:
            try:
                send_telegram(token, chat_id, text)
                self.notify_queue.put((key, price, None))
            except TelegramError as exc:
                self.notify_queue.put((key, price, str(exc)))
            except Exception as exc:  # anything unexpected must still report back, or the poller would wait forever
                self.notify_queue.put((key, price, f"unexpected error: {exc}"))

    def _poll_notify_queue(self):
        try:
            while True:
                key, price, error = self.notify_queue.get_nowait()
                self.notify_outstanding -= 1
                self.pending_alerts.discard(key)
                watch = self.watch_for(*key)
                if error is None:
                    self.watch_errors.pop(key, None)
                    if watch is not None:  # (it may have been removed while the message was in flight)
                        watch["last_alert_price"] = price
                        watch["last_alert_at"] = datetime.now().strftime(HISTORY_TIME_FORMAT)
                        self._persist_settings()
                    self._set_alert_note(f"Alert sent: {key[1]} {fmt_money(price)} at {datetime.now():%H:%M}")
                else:
                    self.watch_errors[key] = error
                    self._set_alert_note(f"Alert FAILED for {key[1]} (see Alerts > View Price Watches)", error=True)
        except queue.Empty:
            pass
        self._after_watches_changed()
        if self.notify_outstanding > 0:
            self.after(200, self._poll_notify_queue)
        else:
            self.notify_polling = False

    def populate_tree(self):
        self.tree.delete(*self.tree.get_children())
        for rank, city, state, code, dist, date_obj, price, mpd in self.display_rows:
            change = self.price_changes.get(code)
            tags = (f"price_{change}",) if change else ()
            delta = self.price_deltas.get(code)
            if delta is None:
                delta_text = ""  # no previous price to compare against
            elif delta == 0:
                delta_text = "$0"
            else:
                delta_text = f"{'+' if delta > 0 else '-'}${abs(delta):,}"
            self.tree.insert(
                "", "end",
                values=(rank, city, state, code, f"{dist:,}", fmt_date(date_obj), f"${price:,}", delta_text, f"{mpd:.2f}"),
                tags=tags,
            )

    def sort_by(self, column):
        """Re-sorts self.display_rows in place (clicking the same header
        again reverses direction). This only changes display/export order
        — each row's Rank value was fixed at fetch time in fetch_deals()
        and is never recalculated here, so it keeps showing each
        destination's true price rank even when the table is sorted by,
        say, Distance instead of Price."""
        if self.sort_column == column:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_column = column
            self.sort_reverse = False

        idx = COLUMN_INDEX.get(column)
        if column == "change":
            # Deltas live outside the row tuples; rows with no previous price sort as 0.
            key = lambda row: self.price_deltas.get(row[3], 0)
        elif column in TEXT_COLUMNS:
            key = lambda row: str(row[idx]).lower()
        else:
            key = lambda row: row[idx]
        self.display_rows.sort(key=key, reverse=self.sort_reverse)

        arrow = " ▼" if self.sort_reverse else " ▲"
        for col, (label, _width) in self.column_labels.items():
            self.tree.heading(col, text=label + (arrow if col == column else ""))

        self.populate_tree()

    def do_export_pdf(self):
        path = output_path(self.origin_code, "pdf")
        try:
            export_pdf(self.display_rows, path, self.origin_code, self.origin_label)
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc))
            return
        self.last_pdf_path = path
        self.open_pdf_btn.config(state="normal")
        self.status_var.set(f"PDF saved to {path}")
        messagebox.showinfo("Export complete", f"PDF saved to:\n{path}")

    def do_export_xlsx(self):
        path = output_path(self.origin_code, "xlsx")
        try:
            export_xlsx(self.display_rows, path, self.origin_code, self.origin_label)
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc))
            return
        self.last_xlsx_path = path
        self.open_xlsx_btn.config(state="normal")
        self.status_var.set(f"XLSX saved to {path}")
        messagebox.showinfo("Export complete", f"XLSX saved to:\n{path}")

    def open_pdf(self):
        os.startfile(self.last_pdf_path)

    def open_xlsx(self):
        os.startfile(self.last_xlsx_path)

    def show_about(self):
        messagebox.showinfo(
            f"About {APP_NAME}",
            f"{APP_NAME}\n"
            f"Version {APP_VERSION}\n\n"
            f"{APP_DESCRIPTION}\n\n"
            f"Contact: {AUTHOR_NAME} <{AUTHOR_EMAIL}>",
        )

    def toggle_auto_refresh(self):
        if self.auto_refresh_active:
            self.stop_auto_refresh()
        else:
            self.start_auto_refresh()

    def start_auto_refresh(self):
        text = self.interval_var.get().strip()
        if not text:
            # Blank box means "leave auto-refresh off" — not an error, just a no-op.
            self.status_var.set("Enter a number of minutes to enable auto-refresh.")
            return

        try:
            minutes = float(text)
        except ValueError:
            messagebox.showerror("Invalid interval", f'"{text}" isn\'t a number.')
            return
        if minutes < MIN_AUTO_REFRESH_MINUTES:
            messagebox.showerror(
                "Invalid interval",
                f"Enter {MIN_AUTO_REFRESH_MINUTES} minute(s) or more.",
            )
            return

        self.auto_refresh_ms = int(minutes * 60 * 1000)
        self.auto_refresh_active = True
        self.interval_entry.config(state="disabled")
        self.auto_refresh_btn.config(text="Stop Auto-Refresh")
        self.status_var.set(f"Auto-refresh on: fetching every {minutes:g} minute(s).")

        # Fetch right away, then keep firing every interval from here on.
        if not self.fetching:
            self.start_fetch(auto=True)
        self.auto_refresh_job = self.after(self.auto_refresh_ms, self._auto_refresh_tick)

    def stop_auto_refresh(self):
        if self.auto_refresh_job is not None:
            self.after_cancel(self.auto_refresh_job)
            self.auto_refresh_job = None
        self._cancel_network_retry()
        self.auto_refresh_active = False
        self.interval_entry.config(state="normal")
        self.auto_refresh_btn.config(text="Start Auto-Refresh")
        if not self.fetching:
            self.status_var.set("Auto-refresh off.")

    def _auto_refresh_tick(self):
        if not self.auto_refresh_active:
            return
        # If a fetch is still running (slow network, etc.) skip this tick
        # rather than piling up overlapping requests — the next one will
        # pick it up in another interval.
        if not self.fetching:
            self.start_fetch(auto=True)
        self.auto_refresh_job = self.after(self.auto_refresh_ms, self._auto_refresh_tick)

    def _on_close(self):
        if self.auto_refresh_job is not None:
            self.after_cancel(self.auto_refresh_job)
        self._cancel_network_retry()
        if self.origins_retry_job is not None:
            self.after_cancel(self.origins_retry_job)
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
